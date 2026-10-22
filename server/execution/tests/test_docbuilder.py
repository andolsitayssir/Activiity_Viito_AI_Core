"""DocBuilder cell/slide script generation + file text extraction (offline).

The live OnlyOffice Document Server runs were verified manually: xlsx row
append ("C, 300" appended as row 4) and pptx text-box insert both succeeded.
This file locks the JS generation + extraction contracts so regressions are
caught without a live container.
"""

from __future__ import annotations

import io


def _build(*args, **kwargs) -> str:
    from server.webapp.docbuilder import build_edit_script

    return build_edit_script(*args, **kwargs)


class TestCellScript:
    def test_append_row_uses_getrange_setvalue(self):
        js = _build(
            "http://x/f.xlsx", "cell", "ajouter une nouvelle ligne C, 300", "out.xlsx",
            modify={"headers": [], "rows_data": [["C", "300"]], "rows_num": 0, "cols_num": 0},
        )
        assert 'builder.OpenFile("http://x/f.xlsx")' in js
        assert "Api.GetActiveSheet()" in js
        assert 'rows = [["C", "300"]]' in js
        assert "GetRange(OCColLetter(cj + 1) + r).SetValue(OCVal(rows[ri][cj])" in js
        assert 'builder.SaveFile("xlsx", "out.xlsx")' in js
        assert not js.count("GetUsedRange")  # not available in the DS build

    def test_append_row_fallback_from_nl_text(self):
        js = _build("http://x/f.xlsx", "cell", "ajouter une nouvelle ligne C, 300", "out.xlsx")
        assert '"C", "300"' in js
        # the "une nouvelle ligne" phrase must be stripped, not a cell value
        assert "nouvelle ligne" not in js

    def test_headers_overwrite_row_one(self):
        js = _build(
            "http://x/f.xlsx", "cell", "", "out.xlsx",
            modify={"headers": ["Produit", "Ventes"], "rows_data": [], "rows_num": 0, "cols_num": 0},
        )
        assert 'oSheet.GetRange(OCColLetter(ci + 1) + "1")' in js
        assert '"Produit"' in js

    def test_insert_table_from_data(self):
        js = _build(
            "http://x/f.xlsx", "cell", "", "out.xlsx",
            table={"headers": ["A", "B"], "rows_data": [["1", "2"]]},
        )
        assert 'oSheet.GetRange("A1").SetValue("A")' in js
        assert 'oSheet.GetRange("B2").SetValue(OCVal("2"))' in js

    def test_insert_uses_verbatim_text_row(self):
        import json as _json

        js = _build(
            "http://x/f.xlsx", "cell", "", "out.xlsx",
            insert={"text": "note de synthèse", "at_beginning": False},
        )
        assert _json.dumps("note de synthèse") in js  # JSON-escaped (non-ascii)


class TestSlideScript:
    def test_textbox_added_via_shape(self):
        js = _build(
            "http://x/f.pptx", "slide", "ajoute du texte", "out.pptx",
            insert={"text": "Hello slide", "at_beginning": False},
        )
        assert "Api.GetPresentation()" in js
        assert "Api.CreateShape(\"rect\", 11000000, 1200000, 500000, 5400000)" in js
        assert 'oParagraph.AddText("Hello slide")' in js
        assert "oSlide.AddObject(oShape)" in js
        assert 'builder.SaveFile("pptx", "out.pptx")' in js

    def test_modify_table_rendered_as_text(self):
        js = _build(
            "http://x/f.pptx", "slide", "", "out.pptx",
            modify={"headers": ["Produit", "Ventes"], "rows_data": [["A", "100"]]},
        )
        assert 'Produit | Ventes' in js
        assert 'A | 100' in js


class TestWordPreserved:
    def test_word_script_unchanged(self):
        js = _build("http://x/f.docx", "word", "ajoute un titre Bonjour", "out.docx")
        assert "var oDocument = Api.GetDocument()" in js
        assert "oDocument.AddElement(0, oParagraph)" in js


class TestTextExtraction:
    def _extract(self, content: bytes, ext: str) -> str | None:
        from server.webapp.onlyoffice import _extract_file_text

        return _extract_file_text(content, ext)

    def test_xlsx_extraction(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Donnees"
        ws.append(["Produit", "Ventes"])
        ws.append(["A", 100])
        ws.append(["B", 200])
        buf = io.BytesIO()
        wb.save(buf)
        text = self._extract(buf.getvalue(), "xlsx")
        assert "Donnees" in text
        assert "Produit | Ventes" in text
        assert "A | 100" in text
        assert "B | 200" in text

    def test_pptx_extraction(self):
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Titre"
        slide.placeholders[1].text = "Contenu"
        buf = io.BytesIO()
        prs.save(buf)
        text = self._extract(buf.getvalue(), "pptx")
        assert "Slide 1" in text
        assert "Titre" in text

    def test_docx_extraction(self):
        import openpyxl
        from docx import Document

        doc = Document()
        doc.add_paragraph("Bonjour le monde")
        buf = io.BytesIO()
        doc.save(buf)
        text = self._extract(buf.getvalue(), "docx")
        assert "Bonjour le monde" in text

    def test_pdf_extraction(self):
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(100, 700, "Rapport PDF test")
        c.showPage()
        c.save()
        text = self._extract(buf.getvalue(), "pdf")
        assert "Rapport PDF test" in text

    def test_binary_returns_none(self):
        assert self._extract(b"\x89PNG\r\n\x1a\n\x00", "png") is None


class TestCellGuard:
    """Idempotency guard: re-running a row-append must not duplicate."""

    _path = "_guard_test.xlsx"

    def _make_sheet(self, rows: list[list]):
        import openpyxl
        from server.execution.executors.file_storage import _full

        path = _full(self._path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Donnees"
        for row in rows:
            ws.append(row)
        wb.save(path)
        path.unlink(missing_ok=True)
        wb.save(path)

    def test_contains_rows_detects_present_and_absent(self):
        from server.execution.executors.onlyoffice import _cell_contains_rows

        self._make_sheet([["Produit", "Ventes"], ["A", 100], ["B", 200]])
        assert _cell_contains_rows(self._path, [["B", "200"]]) is True
        assert _cell_contains_rows(self._path, [["C", "300"]]) is False
        assert _cell_contains_rows(self._path, [["B", "200"], ["C", "300"]]) is False

    def test_contains_rows_after_append(self):
        from server.execution.executors.onlyoffice import _cell_contains_rows

        self._make_sheet([["Produit", "Ventes"], ["A", 100], ["B", 200], ["C", 300]])
        assert _cell_contains_rows(self._path, [["C", "300"]]) is True

    def test_describe_append_reports_row_number(self):
        from server.execution.executors.onlyoffice import _describe_cell_append

        self._make_sheet([["Produit", "Ventes"], ["A", 100], ["B", 200], ["C", 300]])
        desc = _describe_cell_append(self._path, [["C", "300"]])
        assert "Ligne 4 ajoutée" in desc
        assert "C | 300" in desc